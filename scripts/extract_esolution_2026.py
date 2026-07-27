from collections import defaultdict
from datetime import date, datetime
from hashlib import sha1
from json import dump
from pathlib import Path
import re
import sys

from openpyxl import load_workbook


SOURCE = Path(sys.argv[1] if len(sys.argv) > 1 else '/Users/mymac/Downloads/MONITORING SIPRIMA - ELECTRIC SOLUTION UP 3.xlsx')
OUTPUT = Path(sys.argv[2] if len(sys.argv) > 2 else 'storage/app/imports/esolution-2026.json')


def text(value):
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip() or None


def number(value, default=0):
    if value is None or value == '':
        return default
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def date_value(value):
    if isinstance(value, datetime):
        value = value.date()
    if isinstance(value, date) and value.year > 1900:
        return value.isoformat()
    return None


def normal(value):
    return re.sub(r'[^a-z0-9]', '', (text(value) or '').lower())


def meaningful(value):
    return (text(value) or '').lower() not in {'', '-', '0', 'belum ada spk'}


def coordinate(value, minimum, maximum):
    parsed = number(value, None)
    return parsed if parsed is not None and minimum <= parsed <= maximum else None


def link(cell):
    if cell.hyperlink and cell.hyperlink.target:
        return cell.hyperlink.target
    value = text(cell.value)
    return value if value and value.startswith(('https://', 'http://')) else None


def status(value, has_spk):
    normalized = normal(value)
    mapping = {
        'probing': 'probing',
        'penawaran': 'penawaran',
        'spk': 'spk',
        'bast': 'bast',
        'gagalwo': 'gagal_wo',
    }
    return mapping.get(normalized, 'spk' if has_spk else 'probing')


def payment_status(value):
    return 'lunas' if normal(value) == 'paid' else 'belum_ditagih'


workbook = load_workbook(SOURCE, data_only=True, read_only=False)
master = workbook['Recap E-Solution 2026']
offers_sheet = workbook['Recap Penawaran Esol']

offers_by_number = {}
offers_by_customer = defaultdict(list)
for cells in offers_sheet.iter_rows(min_row=3):
    values = [cell.value for cell in cells]
    input_date = date_value(values[2])
    offer_number = text(values[3])
    if not values[6] or not (input_date and input_date.startswith('2026') or '2026' in (offer_number or '')):
        continue

    offer = {
        'no_penawaran': offer_number,
        'pic_penawaran': text(values[29]) or text(values[4]),
        'nilai_jasa_pelanggan': number(values[10]),
        'nilai_material_pelanggan': number(values[9]),
        'nilai_dpp_pelanggan': number(values[11]),
        'nilai_ppn_pelanggan': number(values[12]),
        'nilai_pph_pelanggan': number(values[13]),
        'nilai_dpp_ppn_pelanggan': number(values[11]) + number(values[12]),
        'nilai_terbayar_pelanggan': number(values[14]),
        'nilai_jasa_mitra': number(values[16]),
        'nilai_material_mitra': number(values[15]),
        'nilai_dpp_mitra': number(values[17]),
        'nilai_ppn_mitra': number(values[18]),
        'nilai_pph_mitra': number(values[19]),
        'nilai_dpp_ppn_mitra': number(values[17]) + number(values[18]),
        'nilai_terbayar_mitra': number(values[20]),
        'margin_dpp': number(values[21]),
        'margin_ppn': number(values[22]),
        'margin_persen': number(values[23]),
    }
    if not offer['no_penawaran'] or not offer['pic_penawaran']:
        continue
    offers_by_number[offer['no_penawaran']] = offer
    offers_by_customer[normal(values[6])].append(offer)

groups = defaultdict(list)
for row_number, cells in enumerate(master.iter_rows(min_row=4), start=4):
    if not cells[4].value:
        continue
    spk_number = text(cells[27].value)
    if meaningful(spk_number):
        group_key = f'spk:{normal(spk_number)}'
    else:
        group_key = f'fallback:{normal(cells[4].value)}:{normal(cells[5].value)}:{date_value(cells[18].value) or row_number}'
    groups[group_key].append(cells)

projects = []
for group_key, rows in groups.items():
    first = rows[0]
    source_status = status(first[2].value, meaningful(first[27].value))
    project = {
        'tanggal_probing': date_value(first[18].value),
        'unit_layanan': text(first[16].value),
        'nama_pic_beyond_kwh': None,
        'nama_perusahaan': text(first[4].value),
        'alamat_perusahaan': text(first[10].value) or '-',
        'latitude': coordinate(first[11].value, -90, 90),
        'longitude': coordinate(first[12].value, -180, 180),
        'scope_pekerjaan': 'ev_ecosystem' if 'ev' in normal(first[7].value) else 'electric_services',
        'potensi_pekerjaan': text(first[9].value),
        'nama_pic_pelanggan': text(first[13].value),
        'nomor_pic_pelanggan': text(first[14].value),
        'tanggal_inspeksi': date_value(first[20].value),
        'opsi_selanjutnya': 'sudah_punya_vendor_lain' if source_status == 'gagal_wo' else None,
        'status': source_status,
        'nilai_project_dpp': number(first[24].value, None),
        'nilai_project_inc_ppn': number(first[26].value, None),
    }

    input_offer = text(first[3].value)
    candidates = offers_by_customer[normal(first[4].value)]
    offer = offers_by_number.get(input_offer)
    if offer is None and len(candidates) == 1:
        offer = candidates[0]

    has_historical_relations = any(text(row[37].value) or text(row[56].value) for row in rows)
    spk = None
    spk_number = text(first[27].value)
    if meaningful(spk_number) or has_historical_relations:
        partner = next((text(row[56].value) for row in rows if text(row[56].value)), None)
        spk = {
            'no_spk_pelanggan': spk_number if meaningful(spk_number) else None,
            'tanggal_awal_kontrak': date_value(first[28].value),
            'tanggal_akhir_kontrak': date_value(first[29].value),
            'tanggal_pekerjaan': None,
            'nama_pic_bast': None,
            'nama_pic_bapp': None,
            'nama_mitra': partner or '-',
            'berkas_spk': link(first[51]),
        }

    bast = []
    tagihan_mitra = []
    for row in rows:
        if text(row[37].value):
            bast.append({
                'termin': text(row[32].value) or '-',
                'persentase': number(row[33].value) * 100 if 0 <= number(row[33].value) <= 1 else number(row[33].value),
                'no_bast': text(row[37].value),
                'tanggal_bast': date_value(row[38].value),
                'no_invoice': text(row[41].value),
                'tanggal_invoice': date_value(row[42].value),
                'no_faktur_pajak': text(row[43].value),
                'tanggal_faktur_pajak': date_value(row[44].value),
                'nominal_bayar_pelanggan': number(row[45].value),
                'tanggal_bayar_pelanggan': date_value(row[46].value),
                'status_bayar_pelanggan': payment_status(row[48].value),
                'keterangan': text(row[49].value),
                'berkas_penawaran': link(row[50]),
                'berkas_bast_bapp': link(row[52]),
                'berkas_tagihan': link(row[53]),
            })

        if text(row[56].value):
            tagihan_mitra.append({
                'nama_mitra': text(row[56].value),
                'no_spk': text(row[57].value) or (spk['no_spk_pelanggan'] if spk else None),
                'tanggal_spk': date_value(row[58].value) or (spk['tanggal_awal_kontrak'] if spk else None),
                'no_parking': text(row[59].value),
                'material': number(row[60].value),
                'jasa': number(row[61].value),
                'nilai_spk_exc_ppn': number(row[62].value),
                'nilai_spk_inc_ppn': number(row[63].value),
                'status_smartpayment': text(row[64].value),
                'tgl_verifikator': date_value(row[65].value),
                'tgl_approval': date_value(row[66].value),
                'tgl_man_perbendaharaan': date_value(row[67].value),
                'durasi_lama_hari': int(number(row[68].value)) if row[68].value is not None else None,
                'nilai_terbayar': number(row[69].value),
                'tanggal_bayar': date_value(row[70].value),
                'status_nodin': text(row[71].value),
                'tanggal_nodin': date_value(row[72].value),
                'no_nodin': text(row[73].value),
                'scan_nodin': link(row[74]),
                'scan_tagihan_mitra': link(row[75]),
            })

    source_key = 'recap-2026:' + sha1(group_key.encode()).hexdigest()
    projects.append({
        'source_key': source_key,
        'project': project,
        'penawaran': offer,
        'spk': spk,
        'bast': bast,
        'tagihan_mitra': tagihan_mitra,
    })

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
with OUTPUT.open('w', encoding='utf-8') as handle:
    dump(projects, handle, ensure_ascii=False, indent=2)

print(f'{len(projects)} proyek diekstrak ke {OUTPUT}')
